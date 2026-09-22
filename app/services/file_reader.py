import csv
import io
from openpyxl import load_workbook

def read_csv_bytes(data: bytes):
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]

def read_excel_bytes(data: bytes):
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(x).strip() if x is not None else "" for x in rows[0]]
    output = []
    for values in rows[1:]:
        row = {}
        for i, header in enumerate(headers):
            if header:
                row[header] = values[i] if i < len(values) else None
        if any(v is not None for v in row.values()):
            output.append(row)
    return output
